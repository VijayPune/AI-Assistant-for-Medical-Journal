from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / "team/Clinical_Lead_Job_Description.md"
output = ROOT / "team/Clinical_Lead_Job_Description.docx"
lines = source.read_text(encoding="utf-8").splitlines()

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

NAVY = "17324D"
TEAL = "147D83"
INK = "263746"
PALE = "EAF2F5"
WHITE = "FFFFFF"

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for style_name, size, color in (("Title", 25, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 11.5, TEAL), ("Heading 3", 10.5, NAVY)):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display" if style_name != "Title" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(11 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.add_run("Clinical Lead Job Description")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle.paragraph_format.space_after = Pt(15)
run = subtitle.add_run("Pediatrics and Evidence Quality | Part-time Clinical Adviser")
run.bold = True
run.font.size = Pt(12)
run.font.color.rgb = RGBColor.from_string(TEAL)

def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margin(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def add_inline(paragraph, text: str, *, font_size: float | None = None) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            r = paragraph.add_run(text[cursor:match.start()])
            if font_size: r.font.size = Pt(font_size)
        token = match.group(0)
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2]); r.bold = True
        elif token.startswith("*"):
            r = paragraph.add_run(token[1:-1]); r.italic = True
        else:
            r = paragraph.add_run(token[1:-1]); r.font.name = "Aptos Mono"; r.font.size = Pt(9)
        if font_size: r.font.size = Pt(font_size)
        cursor = match.end()
    if cursor < len(text):
        r = paragraph.add_run(text[cursor:])
        if font_size: r.font.size = Pt(font_size)

def add_table(table_lines: list[str]) -> None:
    rows = []
    for line in table_lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Grid"
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for col_index in range(width):
            cell = cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margin(cell)
            value = row_data[col_index] if col_index < len(row_data) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value, font_size=9)
            if row_index == 0:
                shade(cell, NAVY)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
                    run.bold = True
            elif row_index % 2 == 0:
                shade(cell, PALE)
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

# Start after the Markdown title and subtitle; the Word title is styled independently.
idx = 0
while idx < len(lines) and (not lines[idx].strip() or lines[idx].startswith("# ") or lines[idx].startswith("## ")):
    idx += 1

while idx < len(lines):
    raw = lines[idx].strip()
    if not raw:
        idx += 1
        continue
    if raw.startswith("|"):
        table_lines = []
        while idx < len(lines) and lines[idx].strip().startswith("|"):
            table_lines.append(lines[idx]); idx += 1
        add_table(table_lines)
        continue
    heading = re.match(r"^(#{3,6})\s+(.*)$", raw)
    if heading:
        level = len(heading.group(1))
        style = "Heading 1" if level == 3 else "Heading 2" if level == 4 else "Heading 3"
        p = doc.add_paragraph(style=style)
        if heading.group(2).startswith("11. Suggested interview questions"):
            p.paragraph_format.page_break_before = True
        add_inline(p, heading.group(2))
        idx += 1
        continue
    bullet = re.match(r"^[-*]\s+(.*)$", raw)
    if bullet:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, bullet.group(1))
        idx += 1
        continue
    numbered = re.match(r"^\d+\.\s+(.*)$", raw)
    if numbered:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(4)
        add_inline(p, raw)
        idx += 1
        continue
    p = doc.add_paragraph()
    add_inline(p, raw)
    idx += 1

# Simple footer with automatic page numbering.
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.paragraph_format.space_before = Pt(4)
fr = footer.add_run("Clinical Lead Job Description  |  ")
fr.font.size = Pt(8); fr.font.color.rgb = RGBColor.from_string("667788")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.core_properties.title = "Clinical Lead Job Description"
doc.core_properties.subject = "Pediatrics and evidence quality adviser for the Medical Journal Literature Assistant POC"
doc.core_properties.keywords = "Pediatrics, clinical evaluation, evidence quality, AI literature assistant"
doc.save(output)
print(output)
